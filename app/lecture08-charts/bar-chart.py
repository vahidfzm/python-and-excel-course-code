###############################################################
################# https://www.fardanesh.ir ####################
###############################################################


from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, Series

from copy import deepcopy



wb=load_workbook('lecture08-charts/data1.xlsx')
ws=wb['BarCharts']








chart1 = BarChart()
chart1.type = "col"
data = Reference(ws, min_col=3, min_row=4, max_row=15)
labels = Reference(ws, min_col=2, min_row=5, max_row=15)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(labels)
ws.add_chart(chart1, "E4")








chart2 = BarChart()
chart2.type = "col"
chart2.style = 10
chart2.title = "Seasonal sales Bar Chart"
chart2.y_axis.title = 'Sale'
chart2.x_axis.title = 'Year'
data = Reference(ws, min_col=3,max_col=6, min_row=22, max_row=33)
cats = Reference(ws, min_col=2, min_row=23, max_row=33)
chart2.add_data(data, titles_from_data=True)
chart2.set_categories(cats)
chart2.shape = 4
ws.add_chart(chart2, "H22")








chart2Stacked=deepcopy(chart2)
chart2Stacked.grouping = "stacked"
chart2Stacked.overlap = 100
ws.add_chart(chart2Stacked, "R22")







wb.save(filename='lecture08-charts/data1WithChart.xlsx')
