###############################################################
################# https://www.fardanesh.ir ####################
###############################################################


from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


wb=load_workbook('lecture09-tables-images/list1.xlsx')
ws1 = wb['Sheet1']
ws3 = wb['Sheet3']


tbl1=Table(displayName="Table3",ref="A1:H20")


# TableStyleInfo
# ["TableStyleMedium{0}".format(i) for i in range(1, 29)]
# ["TableStyleLight{0}".format(i) for i in range(1, 22)]
# ["TableStyleDark{0}".format(i) for i in range(1, 12)]
tableStyle1 = TableStyleInfo(name="TableStyleMedium18", showFirstColumn=False,showLastColumn=False, showRowStripes=True, showColumnStripes=True)
tableStyle2 = TableStyleInfo(name="TableStyleLight3", showFirstColumn=False,showLastColumn=False, showRowStripes=True, showColumnStripes=True)

tbl1.tableStyleInfo=tableStyle1

ws1.add_table(tbl1)






ws3.tables['Table1'].tableStyleInfo=tableStyle2




for tbl in ws3.tables.items():
    print(tbl)


for tableName,tableRange in ws3.tables.items():
    print(tableName)
    print(tableRange)


del ws3.tables['Table2']


wb.save(filename='lecture09-tables-images/list1WithTable.xlsx')
